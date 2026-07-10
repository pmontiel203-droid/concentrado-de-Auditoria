"""Tests for pre-assembled concentrado (SEMANA sheets) import feature."""
import io
import os
import pytest
import requests
import openpyxl

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
API = f"{BASE_URL}/api"
SAMPLE = "/app/tests/samples/concentrado_armado_junio.xlsx"

MES = 6
ANIO = 2026
SEMANAS_ESPERADAS = {22, 23, 24, 25}
TIENDAS_ESPERADAS = 17


@pytest.fixture(scope="module")
def sample_bytes():
    with open(SAMPLE, "rb") as f:
        return f.read()


@pytest.fixture(scope="module", autouse=True)
def clean_state():
    for s in SEMANAS_ESPERADAS:
        requests.delete(f"{API}/registros", params={"semana": s, "anio": ANIO}, timeout=30)
    yield


def _post_import(payload_bytes, filename="concentrado_armado_junio.xlsx"):
    files = {"files": (filename, payload_bytes,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    return requests.post(f"{API}/import", files=files, timeout=120)


# ---------- 1. Import pre-assembled file ----------
def test_01_import_concentrado_armado(sample_bytes):
    r = _post_import(sample_bytes)
    assert r.status_code == 200, r.text
    resultados = r.json()["resultados"]
    assert len(resultados) == 1
    e = resultados[0]
    assert e["resultado"] == "IMPORTADO", e
    # ~882 registros
    assert 800 <= e["registros"] <= 950, f"registros out of expected range: {e['registros']}"
    # tienda field should be summary
    assert "tienda" in e["tienda"] or e["tienda"].endswith("tiendas"), e
    assert str(TIENDAS_ESPERADAS) in e["tienda"], e
    # semanas string contains 22..25
    for s in SEMANAS_ESPERADAS:
        assert str(s) in str(e["semana"]), f"semana {s} missing in {e['semana']}"


# ---------- 2. /api/semanas populated ----------
def test_02_semanas_endpoint():
    data = requests.get(f"{API}/semanas", timeout=30).json()
    for s in SEMANAS_ESPERADAS:
        matches = [row for row in data if row["semana"] == s and row["anio"] == ANIO]
        assert matches, f"semana {s}/{ANIO} not in /api/semanas"
        row = matches[0]
        assert row["mes"] == MES, f"mes mismatch for semana {s}: {row}"
        assert row["num_tiendas"] == TIENDAS_ESPERADAS, f"num_tiendas={row['num_tiendas']} for semana {s}"


# ---------- 3. /api/registros fields ----------
def test_03_registros_semana_25():
    regs = requests.get(f"{API}/registros",
                        params={"semana": 25, "anio": ANIO}, timeout=30).json()
    assert len(regs) > 0
    for field in ["detalle_full", "area", "catalogo", "responsable", "tienda",
                  "semana", "mes", "anio", "auditados", "negado", "recuperado"]:
        assert field in regs[0], f"missing field {field} in registro"
    # recuperado should be 0 for pre-assembled source
    total_rec = sum(r["recuperado"] for r in regs)
    assert total_rec == 0, f"expected recuperado=0 for pre-assembled, got {total_rec}"
    # detalle_full is a dict with '3'..'37'
    df = regs[0]["detalle_full"]
    assert isinstance(df, dict)
    for c in range(3, 38):
        assert str(c) in df, f"detalle_full missing col {c}"


# ---------- 4. Idempotency / overwrite ----------
def test_04_reimport_idempotent(sample_bytes):
    # count before
    def count_all():
        total = 0
        for s in SEMANAS_ESPERADAS:
            total += len(requests.get(f"{API}/registros",
                                      params={"semana": s, "anio": ANIO},
                                      timeout=30).json())
        return total

    before = count_all()
    assert before > 0
    r = _post_import(sample_bytes)
    assert r.status_code == 200
    e = r.json()["resultados"][0]
    assert e["resultado"] == "IMPORTADO"
    after = count_all()
    assert before == after, f"Idempotency failed: before={before}, after={after}"

    # still 4 weeks x 17 tiendas
    for s in SEMANAS_ESPERADAS:
        regs = requests.get(f"{API}/registros",
                            params={"semana": s, "anio": ANIO}, timeout=30).json()
        tiendas = {r["tienda"] for r in regs}
        assert len(tiendas) == TIENDAS_ESPERADAS, \
            f"semana {s}: expected {TIENDAS_ESPERADAS} tiendas, got {len(tiendas)}"


# ---------- 5. Dashboard for month 6/2026 ----------
def test_05_dashboard_junio():
    d = requests.get(f"{API}/dashboard",
                     params={"mes": MES, "anio": ANIO}, timeout=30).json()
    assert d["tiendas_recibidas"] == TIENDAS_ESPERADAS, d
    assert "pct_negado" in d and isinstance(d["pct_negado"], (int, float))
    assert d["ranking_tiendas"] and len(d["ranking_tiendas"]) >= 1
    assert d["por_catalogo"] and len(d["por_catalogo"]) >= 1


# ---------- 6. Reporte mensual xlsx ----------
def test_06_reporte_mensual_junio():
    r = requests.get(f"{API}/reporte/mensual",
                     params={"mes": MES, "anio": ANIO}, timeout=90)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "SEMANA" in wb.sheetnames
    assert "Concentrado " in wb.sheetnames
    ws = wb["SEMANA"]

    # Row 7 Aguascalientes col3>0 (Calzado area)
    v_name = ws.cell(row=7, column=1).value
    assert v_name and "aguascalientes" in str(v_name).lower(), \
        f"row 7 col1 expected Aguascalientes, got {v_name}"
    col3 = ws.cell(row=7, column=3).value
    assert col3 is not None and float(col3) > 0, f"row 7 col3 empty: {col3}"

    # Concentrado sheet Aguascalientes col6>0
    ws2 = wb["Concentrado "]
    found = False
    for row in range(1, ws2.max_row + 1):
        v = ws2.cell(row=row, column=2).value
        if v and "aguascalientes" in str(v).lower():
            c6 = ws2.cell(row=row, column=6).value
            assert c6 and float(c6) > 0, f"Concentrado Aguascalientes col6 empty: {c6}"
            found = True
            break
    assert found, "Aguascalientes row not found in Concentrado sheet"


# ---------- 7. Reporte anual xlsx ----------
def test_07_reporte_anual():
    r = requests.get(f"{API}/reporte/anual", params={"anio": ANIO}, timeout=90)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Concentrado " in wb.sheetnames


# ---------- 8. Regression: per-store file /app/tests/samples/checklist_sem26.xlsx ----------
def test_08_regression_per_store_sin_forzar():
    # Clean semana 26 first
    requests.delete(f"{API}/registros", params={"semana": 26, "anio": ANIO}, timeout=30)
    with open("/app/tests/samples/checklist_sem26.xlsx", "rb") as f:
        b = f.read()
    files = {"files": ("checklist_sem26.xlsx", b,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{API}/import", files=files, timeout=60)
    assert r.status_code == 200
    e = r.json()["resultados"][0]
    assert e["resultado"] == "CON DIFERENCIAS", e


def test_09_regression_per_store_forzar():
    with open("/app/tests/samples/checklist_sem26.xlsx", "rb") as f:
        b = f.read()
    files = {"files": ("checklist_sem26.xlsx", b,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"forzar": "true"}
    r = requests.post(f"{API}/import", files=files, data=data, timeout=60)
    assert r.status_code == 200
    e = r.json()["resultados"][0]
    assert e["resultado"] == "IMPORTADO (FORZADO)", e
    assert e["registros"] == 16
    assert "ECATEPEC" in e["tienda"].upper()
    assert e["semana"] == 26


# ---------- 10. OMITIDO for xlsx with neither sheets nor SEMANA ----------
def test_10_omitido_no_sheets():
    wb = openpyxl.Workbook()
    wb.active.title = "Random"
    wb.active["A1"] = "hola"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    files = {"files": ("bad.xlsx", buf.read(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{API}/import", files=files, timeout=30)
    assert r.status_code == 200
    e = r.json()["resultados"][0]
    assert e["resultado"] == "OMITIDO", e
