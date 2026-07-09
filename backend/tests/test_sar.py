"""SAR backend integration tests."""
import io
import os
import pytest
import requests
import openpyxl

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://audit-recovery.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"
SAMPLE = "/app/tests/samples/checklist_sem26.xlsx"

SEMANA = 26
ANIO = 2026


@pytest.fixture(scope="module")
def sample_bytes():
    with open(SAMPLE, "rb") as f:
        return f.read()


@pytest.fixture(scope="module", autouse=True)
def clean_state():
    # Ensure clean DB for semana 26 before test run
    requests.delete(f"{API}/registros", params={"semana": SEMANA, "anio": ANIO}, timeout=30)
    yield
    # optional cleanup after
    requests.delete(f"{API}/registros", params={"semana": SEMANA, "anio": ANIO}, timeout=30)


def _post_import(sample_bytes, forzar=None, filename="checklist_sem26.xlsx"):
    files = {"files": (filename, sample_bytes,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {}
    if forzar is not None:
        data["forzar"] = str(forzar).lower()
    return requests.post(f"{API}/import", files=files, data=data, timeout=60)


# ---------- 1. No-forzar => CON DIFERENCIAS, no records ----------
def test_import_sin_forzar_con_diferencias(sample_bytes):
    r = _post_import(sample_bytes, forzar=None)
    assert r.status_code == 200, r.text
    resultados = r.json()["resultados"]
    assert len(resultados) == 1
    entry = resultados[0]
    assert entry["resultado"] == "CON DIFERENCIAS", entry
    assert entry["tienda"], "should have tienda parsed"
    assert entry["semana"] == SEMANA
    v = entry.get("validacion", {})
    assert v.get("coincide") is False
    assert v.get("diferencia", 0) != 0

    # Verify no records were stored
    regs = requests.get(f"{API}/registros", params={"semana": SEMANA, "anio": ANIO}, timeout=30).json()
    assert regs == [], f"Expected no records but got {len(regs)}"


# ---------- 2. Forzar => IMPORTADO (FORZADO), 16 records ----------
def test_import_forzar_true(sample_bytes):
    r = _post_import(sample_bytes, forzar=True)
    assert r.status_code == 200, r.text
    entry = r.json()["resultados"][0]
    assert entry["resultado"] == "IMPORTADO (FORZADO)", entry
    assert entry["registros"] == 16
    assert "ECATEPEC" in entry["tienda"].upper()
    assert entry["semana"] == SEMANA


# ---------- 3. Idempotency ----------
def test_import_idempotent(sample_bytes):
    _post_import(sample_bytes, forzar=True)
    _post_import(sample_bytes, forzar=True)
    regs = requests.get(f"{API}/registros", params={"semana": SEMANA, "anio": ANIO}, timeout=30).json()
    # Only Ecatepec has 16 records
    ecatepec = [r for r in regs if "ECATEPEC" in (r.get("tienda") or "").upper()]
    assert len(ecatepec) == 16, f"Idempotency failed: got {len(ecatepec)} records"


# ---------- 4. Non-audit file => OMITIDO ----------
def test_import_omitido_no_sheets():
    wb = openpyxl.Workbook()
    wb.active.title = "Random"
    wb.active["A1"] = "hola"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    files = {"files": ("bad.xlsx", buf.read(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{API}/import", files=files, timeout=30)
    assert r.status_code == 200
    entry = r.json()["resultados"][0]
    assert entry["resultado"] == "OMITIDO", entry


# ---------- 5. Semanas ----------
def test_get_semanas():
    r = requests.get(f"{API}/semanas", timeout=30)
    assert r.status_code == 200
    data = r.json()
    assert any(s["semana"] == SEMANA and s["anio"] == ANIO for s in data), data


# ---------- 6. Registros field structure ----------
def test_registros_fields():
    regs = requests.get(f"{API}/registros", params={"semana": SEMANA, "anio": ANIO}, timeout=30).json()
    ecatepec = [r for r in regs if "ECATEPEC" in (r.get("tienda") or "").upper()]
    assert len(ecatepec) == 16
    required = ["semana", "mes", "anio", "fecha_importacion", "nombre_archivo",
                "tienda", "responsable", "area", "catalogo", "auditados", "negado",
                "recuperado", "sin_exhibicion", "exhibicion", "frente", "bodega",
                "costo", "observaciones"]
    for field in required:
        assert field in ecatepec[0], f"missing field {field}"


# ---------- 7. Dashboard ----------
def test_dashboard():
    r = requests.get(f"{API}/dashboard", params={"semana": SEMANA, "anio": ANIO}, timeout=30)
    assert r.status_code == 200
    d = r.json()
    for key in ["tiendas_recibidas", "tiendas_pendientes", "pct_negado",
                "pct_recuperado", "ranking_tiendas", "ranking_responsables",
                "por_catalogo", "total_costo"]:
        assert key in d, f"missing {key}"
    assert d["tiendas_pendientes"] == 17 - d["tiendas_recibidas"]
    assert d["tiendas_recibidas"] >= 1


# ---------- 8. Comparativo ----------
def test_comparativo():
    r = requests.get(f"{API}/comparativo", params={"anio": ANIO, "tipo": "semana"}, timeout=30)
    assert r.status_code == 200
    data = r.json()
    assert isinstance(data, list) and len(data) > 0
    row = data[0]
    for k in ["periodo", "pct_negado", "pct_recuperado"]:
        assert k in row


# ---------- 9. Imports log ----------
def test_imports_log():
    r = requests.get(f"{API}/imports", timeout=30)
    assert r.status_code == 200
    logs = r.json()
    assert len(logs) > 0
    entry = logs[0]
    for k in ["fecha", "hora", "semana", "tienda", "archivo", "resultado",
              "observaciones", "usuario"]:
        assert k in entry, f"missing {k}"


# ---------- 10. Reporte semanal excel ----------
def test_reporte_semanal_excel():
    r = requests.get(f"{API}/reporte/semanal", params={"semana": SEMANA, "anio": ANIO}, timeout=60)
    assert r.status_code == 200
    assert "spreadsheet" in r.headers.get("Content-Type", "")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Concentrado " in wb.sheetnames
    ws = wb["Concentrado "]
    # find Ecatepec row in column 2
    found_row = None
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row=row, column=2).value
        if v and "ecatepec" in str(v).lower():
            found_row = row
            break
    assert found_row, "Ecatepec row not found"
    col6 = ws.cell(row=found_row, column=6).value
    col64 = ws.cell(row=found_row, column=64).value
    assert col6 and float(col6) > 0, f"col 6 (auditados) empty at row {found_row}: {col6}"
    assert col64 is not None, f"col 64 empty at row {found_row}"


# ---------- 11. Reporte mensual ----------
def test_reporte_mensual():
    r = requests.get(f"{API}/reporte/mensual", params={"mes": 7, "anio": ANIO}, timeout=60)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Concentrado " in wb.sheetnames


# ---------- 12. Reporte anual ----------
def test_reporte_anual():
    r = requests.get(f"{API}/reporte/anual", params={"anio": ANIO}, timeout=60)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Concentrado " in wb.sheetnames


# ---------- 13. Reporte semanal 404 ----------
def test_reporte_semanal_404():
    r = requests.get(f"{API}/reporte/semanal", params={"semana": 999}, timeout=30)
    assert r.status_code == 404


# ---------- 14. Delete registros ----------
def test_delete_registros():
    # ensure data exists
    with open(SAMPLE, "rb") as f:
        _post_import(f.read(), forzar=True)
    r = requests.delete(f"{API}/registros", params={"semana": SEMANA, "anio": ANIO}, timeout=30)
    assert r.status_code == 200
    body = r.json()
    assert "eliminados" in body
    assert body["eliminados"] >= 16
    # verify gone
    regs = requests.get(f"{API}/registros", params={"semana": SEMANA, "anio": ANIO}, timeout=30).json()
    assert regs == []
