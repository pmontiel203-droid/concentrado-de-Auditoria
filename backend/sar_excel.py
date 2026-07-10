"""SAR - Motor de parseo y generación de reportes Excel."""
import io
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

TEMPLATE_PATH = Path(__file__).parent / "templates" / "machote.xlsx"

REQUIRED_SHEETS = ["Conc Dictamenes", "Exhibiciones"]

# Causas de negado en el orden de "Conc Dictamenes" (col 6..19) que mapean 1:1
# a los grupos de columnas del MACHOTE (col 8,12,16,...,60).
NEGADO_CAUSAS = [
    "Frenteos", "Revueltos", "Bodega Venta", "Bodega Tapanco", "Doble Ubicacion",
    "Excedente", "Origen", "Cambios / Probadores", "Negados sin Recuperar",
    "Exhibiciones", "Bodega Aclaracion", "Aduanas", "Golden o apartados",
    "Control de calidad",
]

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

ROPA_CATALOGOS = {"ROPA DOBLADA", "ROPA COLGADA", "JEANS"}


def norm(s):
    if s is None:
        return ""
    s = str(s).strip()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return s.upper().strip()


def num(v):
    """Coerce a cell value to float; '-', None, text -> 0."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "#VALUE!", "#DIV/0!"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def area_de_catalogo(catalogo):
    return "Ropa" if norm(catalogo) in ROPA_CATALOGOS else "Calzado"


def find_sheet(wb, target):
    tnorm = norm(target)
    for name in wb.sheetnames:
        if norm(name) == tnorm:
            return name
    return None


def es_archivo_valido(wb):
    return all(find_sheet(wb, s) is not None for s in REQUIRED_SHEETS)


def parse_semana(raw):
    m = re.search(r"(\d+)", str(raw or ""))
    return int(m.group(1)) if m else None


def parse_periodo(texto):
    """Extrae mes y año de textos como 'Del 29 al 03 de Julio del 2026'."""
    t = norm(texto)
    mes = None
    for i, nombre in enumerate(MESES, start=1):
        if norm(nombre) in t:
            mes = i
            break
    ym = re.search(r"(20\d{2})", str(texto or ""))
    anio = int(ym.group(1)) if ym else None
    return mes, anio


def parse_conc_dictamenes(wb):
    """Devuelve (meta, registros[]) desde la hoja Conc Dictamenes."""
    ws = wb[find_sheet(wb, "Conc Dictamenes")]

    def cell(r, c):
        return ws.cell(row=r, column=c).value

    tienda = cell(4, 3)
    semana = parse_semana(cell(4, 10))
    periodo = cell(4, 26)
    mes, anio = parse_periodo(periodo)

    # bloque de recuperadas: filas 29..44, catalogo en col 4, causas cols 6..19
    recuperadas = {}
    for r in range(29, 45):
        cat = cell(r, 4)
        if cat is None:
            continue
        recuperadas[norm(cat)] = {
            NEGADO_CAUSAS[i]: num(cell(r, 6 + i)) for i in range(14)
        }

    registros = []
    for r in range(8, 24):
        cat = cell(r, 1)
        if cat is None or norm(cat) in ("", "TOTAL"):
            continue
        catnorm = norm(cat)
        negado_detalle = {NEGADO_CAUSAS[i]: num(cell(r, 6 + i)) for i in range(14)}
        rec_detalle = recuperadas.get(catnorm, {c: 0.0 for c in NEGADO_CAUSAS})
        negado_total = sum(negado_detalle.values())
        recuperado_total = sum(rec_detalle.values())
        # fila completa (col 3..37) para reconstruir la hoja SEMANA con todo el detalle
        detalle_full = {str(c): num(cell(r, c)) for c in range(3, 38)}
        reg = {
            "catalogo": str(cat).strip(),
            "responsable": str(cell(r, 2)).strip() if cell(r, 2) else "",
            "area": area_de_catalogo(cat),
            "total_pasillo": num(cell(r, 3)),
            "auditados": num(cell(r, 4)),
            "pct_cobertura": num(cell(r, 5)),
            "frente": num(cell(r, 6)),
            "bodega": num(cell(r, 8)),
            "negado": negado_total,
            "recuperado": recuperado_total,
            "exhibicion": num(cell(r, 15)),
            "sin_exhibicion": num(cell(r, 21)),
            "costo": num(cell(r, 20)),
            "observaciones": "",
            "negado_detalle": negado_detalle,
            "recuperado_detalle": rec_detalle,
            "detalle_full": detalle_full,
        }
        registros.append(reg)

    meta = {"tienda": str(tienda).strip() if tienda else "", "semana": semana,
            "mes": mes, "anio": anio, "periodo": str(periodo).strip() if periodo else ""}
    return meta, registros


def validar_exhibiciones(wb, registros):
    """Compara totales de Exhibiciones vs Conc Dictamenes. Devuelve dict."""
    ws = wb[find_sheet(wb, "Exhibiciones")]
    total_oportunidades = 0.0
    total_recuperadas = 0.0
    for r in range(1, ws.max_row + 1):
        c3 = ws.cell(row=r, column=3).value
        if c3 and norm(c3).startswith("TOTAL"):
            total_oportunidades += num(ws.cell(row=r, column=19).value)
            total_recuperadas += num(ws.cell(row=r, column=20).value)

    conc_incidencias = sum(
        reg["exhibicion"] + reg["sin_exhibicion"] for reg in registros
    )
    diferencia = round(total_oportunidades - conc_incidencias, 2)
    return {
        "exhibiciones_oportunidades": total_oportunidades,
        "exhibiciones_recuperadas": total_recuperadas,
        "conc_incidencias_exhibicion": conc_incidencias,
        "diferencia": diferencia,
        "coincide": abs(diferencia) < 0.001,
    }


# ---------------- Importación de Concentrados YA ARMADOS ----------------
# Estos archivos (Concentrado de Auditoría y Recuperación) traen varias hojas
# 'SEMANA xx' con el detalle por catálogo y por tienda (Parte 3). Se leen tal cual
# hacia la Base Histórica, SIN validación de Exhibiciones.

def es_concentrado_armado(wb):
    """True si el archivo tiene al menos una hoja 'SEMANA' con bloques por catálogo."""
    for name in wb.sheetnames:
        if norm(name).startswith("SEMANA"):
            ws = wb[name]
            for r in range(1, min(ws.max_row, 130) + 1):
                if norm(ws.cell(row=r, column=2).value) == "RESPONSABLE DE CATALOGO":
                    return True
    return False


def _registro_desde_detalle(ws, row, tienda, semana, mes, anio, periodo):
    cat = ws.cell(row=row, column=1).value
    if cat is None or norm(cat) in ("", "TOTAL"):
        return None
    resp = ws.cell(row=row, column=2).value
    resp = "" if resp is None or str(resp).strip() in ("", "0") else str(resp).strip()
    det = {str(c): num(ws.cell(row=row, column=c).value) for c in range(3, 38)}
    negado_detalle = {NEGADO_CAUSAS[i]: num(ws.cell(row=row, column=6 + i).value) for i in range(14)}
    negado_total = sum(negado_detalle.values())
    # descartar filas totalmente vacías (catálogos no auditados)
    if det["3"] == 0 and det["4"] == 0 and negado_total == 0 and det["20"] == 0:
        return None
    return {
        "catalogo": str(cat).strip(), "responsable": resp,
        "area": area_de_catalogo(cat),
        "total_pasillo": det["3"], "auditados": det["4"], "pct_cobertura": det["5"],
        "frente": det["6"], "bodega": det["8"],
        "negado": negado_total, "recuperado": 0.0,
        "exhibicion": det["15"], "sin_exhibicion": det["21"], "costo": det["20"],
        "observaciones": "",
        "negado_detalle": negado_detalle,
        "recuperado_detalle": {c: 0.0 for c in NEGADO_CAUSAS},
        "detalle_full": det,
        "semana": semana, "mes": mes, "anio": anio,
        "periodo": str(periodo).strip() if periodo else "",
        "tienda": str(tienda).strip(),
    }


def _parse_hoja_semana(ws, semana_defecto):
    sem = parse_semana(ws.cell(row=3, column=10).value) or semana_defecto
    periodo = ws.cell(row=3, column=25).value
    mes, anio = parse_periodo(periodo)
    registros = []
    maxr = min(ws.max_row, 700)
    r = 1
    while r <= maxr:
        if norm(ws.cell(row=r, column=2).value) == "RESPONSABLE DE CATALOGO" and ws.cell(row=r, column=1).value:
            tienda = str(ws.cell(row=r, column=1).value).strip()
            rr = r + 1
            leidas = 0
            while rr <= maxr and leidas < 20:
                if norm(ws.cell(row=rr, column=2).value) == "TOTAL":
                    break
                reg = _registro_desde_detalle(ws, rr, tienda, sem, mes, anio, periodo)
                if reg:
                    registros.append(reg)
                rr += 1
                leidas += 1
            r = rr
        r += 1
    return registros


def parse_concentrado_armado(wb):
    """Devuelve lista plana de registros leídos de todas las hojas SEMANA."""
    out = []
    for name in wb.sheetnames:
        if norm(name).startswith("SEMANA"):
            out.extend(_parse_hoja_semana(wb[name], parse_semana(name)))
    return out



# ---------------- Generación de reportes ----------------

CATALOGO_ORDEN = ["BOTAS", "URBANO", "ESCOLAR", "SANDALIAS", "CONFORT",
                  "VESTIR CASUAL", "INFANTILES", "CABALLEROS", "CHANCLERO",
                  "ROPA DOBLADA", "ROPA COLGADA", "IMPORTADOS", "BARRA DE BEBES",
                  "JEANS", "MOCHILAS", "HASTA"]

# columnas de detalle (3..37); la 5 es % de cobertura (se recalcula en agregados)
DET_COLS = [c for c in range(3, 38)]


def _store_row_map(ws, r_ini, r_fin):
    """Mapea nombre normalizado de tienda -> fila, en un rango de la plantilla."""
    m = {}
    for r in range(r_ini, r_fin + 1):
        v = ws.cell(row=r, column=2).value
        if v and isinstance(v, str) and v.strip():
            m[norm(v)] = r
    return m


def _fill_section(ws, r_ini, r_fin, por_tienda):
    """Rellena una sección (Calzado/Ropa) de la hoja Concentrado del MACHOTE."""
    row_map = _store_row_map(ws, r_ini, r_fin)
    for tienda_norm, agg in por_tienda.items():
        row = row_map.get(tienda_norm)
        if not row:
            continue
        ws.cell(row=row, column=5, value=round(agg["total_pasillo"]))
        ws.cell(row=row, column=6, value=round(agg["auditados"]))
        for i, causa in enumerate(NEGADO_CAUSAS):
            ws.cell(row=row, column=8 + 4 * i, value=round(agg["negado_detalle"].get(causa, 0)))
            ws.cell(row=row, column=10 + 4 * i, value=round(agg["recuperado_detalle"].get(causa, 0)))
        ws.cell(row=row, column=64, value=round(agg["negado"]))
        ws.cell(row=row, column=66, value=round(agg["recuperado"]))
        ws.cell(row=row, column=68, value=round(agg["costo"]))


def _aggregate(registros):
    """Agrupa registros por (area, tienda) para la hoja Concentrado."""
    out = {"Calzado": {}, "Ropa": {}}
    for reg in registros:
        area = reg.get("area", "Calzado")
        area = area if area in out else "Calzado"
        t = norm(reg["tienda"])
        d = out[area].setdefault(t, {
            "total_pasillo": 0.0, "auditados": 0.0, "negado": 0.0,
            "recuperado": 0.0, "costo": 0.0,
            "negado_detalle": {c: 0.0 for c in NEGADO_CAUSAS},
            "recuperado_detalle": {c: 0.0 for c in NEGADO_CAUSAS},
        })
        d["total_pasillo"] += reg["total_pasillo"]
        d["auditados"] += reg["auditados"]
        d["negado"] += reg["negado"]
        d["recuperado"] += reg["recuperado"]
        d["costo"] += reg["costo"]
        for c in NEGADO_CAUSAS:
            d["negado_detalle"][c] += reg.get("negado_detalle", {}).get(c, 0)
            d["recuperado_detalle"][c] += reg.get("recuperado_detalle", {}).get(c, 0)
    return out


def _det(reg, col):
    return (reg.get("detalle_full") or {}).get(str(col), 0.0)


def _fill_semana(ws, registros):
    """Rellena la hoja SEMANA: Calzado por tienda, Ropa por tienda y
    el detalle por catálogo de cada tienda (nivel catálogos)."""
    # localizar cabeceras dinámicamente
    calzado_hdr = ropa_hdr = None
    bloques = {}  # tienda_norm -> fila cabecera del bloque de detalle
    for r in range(1, min(ws.max_row, 500) + 1):
        c1 = ws.cell(row=r, column=1).value
        c2 = ws.cell(row=r, column=2).value
        n2 = norm(c2)
        if n2 == "CALZADO":
            calzado_hdr = r
        elif n2 == "ROPA":
            ropa_hdr = r
        elif n2 == "RESPONSABLE DE CATALOGO" and c1:
            bloques[norm(c1)] = r

    # --- Partes 1 y 2: agregados por tienda (Calzado / Ropa) ---
    def fill_area(hdr, area):
        if not hdr:
            return
        rows = {}
        for r in range(hdr + 1, hdr + 18):
            v = ws.cell(row=r, column=1).value
            if v and isinstance(v, str) and norm(v) not in ("", "TOTAL"):
                rows[norm(v)] = r
        # sumar por tienda los catálogos de esa área
        agg = {}
        for reg in registros:
            if reg.get("area") != area:
                continue
            t = norm(reg["tienda"])
            d = agg.setdefault(t, {c: 0.0 for c in DET_COLS})
            for c in DET_COLS:
                if c == 5:
                    continue
                d[c] += _det(reg, c)
        for tnorm, d in agg.items():
            row = _match_row(rows, tnorm)
            if not row:
                continue
            for c in DET_COLS:
                if c == 5:
                    pas, mod = d.get(3, 0), d.get(4, 0)
                    ws.cell(row=row, column=5, value=(mod / pas) if pas else 0)
                else:
                    ws.cell(row=row, column=c, value=round(d[c], 2))

    fill_area(calzado_hdr, "Calzado")
    fill_area(ropa_hdr, "Ropa")

    # --- Parte 3: detalle por catálogo de cada tienda ---
    por_tienda = {}
    for reg in registros:
        por_tienda.setdefault(norm(reg["tienda"]), {})[norm(reg["catalogo"])] = reg
    for tnorm, hdr in bloques.items():
        cats = _match_dict(por_tienda, tnorm)
        if not cats:
            continue
        for i, cat_nombre in enumerate(CATALOGO_ORDEN):
            row = hdr + 1 + i
            reg = cats.get(norm(cat_nombre))
            ws.cell(row=row, column=1, value=cat_nombre)
            if reg:
                ws.cell(row=row, column=2, value=reg.get("responsable", ""))
                for c in DET_COLS:
                    ws.cell(row=row, column=c, value=round(_det(reg, c), 2))


def _match_row(rows_map, tnorm):
    if tnorm in rows_map:
        return rows_map[tnorm]
    for k, v in rows_map.items():
        if tnorm in k or k in tnorm:
            return v
    return None


def _match_dict(por_tienda, tnorm):
    if tnorm in por_tienda:
        return por_tienda[tnorm]
    for k, v in por_tienda.items():
        if tnorm in k or k in tnorm:
            return v
    return None


def generar_concentrado(registros, etiqueta):
    """Genera un xlsx con formato MACHOTE (hojas SEMANA + Concentrado). Devuelve bytes."""
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    # Hoja de detalle SEMANA (por tienda y por catálogo)
    if "SEMANA" in wb.sheetnames:
        _fill_semana(wb["SEMANA"], registros)
        try:
            wb["SEMANA"].cell(row=3, column=10, value=etiqueta)
        except Exception:
            pass
    # Hoja de presentación Concentrado
    ws = wb["Concentrado "]
    agg = _aggregate(registros)
    _fill_section(ws, 8, 24, agg["Calzado"])
    _fill_section(ws, 39, 55, agg["Ropa"])
    for (r, c) in [(3, 24), (3, 88), (3, 117), (34, 24), (34, 117)]:
        try:
            ws.cell(row=r, column=c, value=etiqueta)
        except Exception:
            pass
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
